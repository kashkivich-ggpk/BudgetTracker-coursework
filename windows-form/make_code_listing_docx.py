from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"C:\Users\user\Desktop\windows-form")
SOURCE_DIR = ROOT / "BudgetTracker"
OUTPUT = ROOT / "Листинг_исходного_кода_BudgetTracker_3_колонки.docx"

SOURCE_FILES = [
    "BudgetService.h",
    "ClientIdentityService.h",
    "DomainModels.h",
    "FinanceService.h",
    "ForecastService.h",
    "IDataStorage.h",
    "InputValidator.h",
    "MyForm.cpp",
    "MyForm.h",
    "PostgreSqlStorage.h",
    "RecurringPaymentService.h",
    "ReportService.h",
    "StorageFactory.h",
    "SupabaseStorage.h",
]


def set_columns(section, count=3, spacing_twips=280):
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(spacing_twips))
    cols.set(qn("w:equalWidth"), "1")
    cols.set(qn("w:sep"), "0")


def set_run_font(run, size=5.5):
    run.font.name = "Courier New"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(size)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Cm(29.7)
section.page_height = Cm(21.0)
section.top_margin = Cm(0.65)
section.bottom_margin = Cm(0.65)
section.left_margin = Cm(0.65)
section.right_margin = Cm(0.65)
section.header_distance = Cm(0.3)
section.footer_distance = Cm(0.3)
set_columns(section, count=3, spacing_twips=280)

normal = doc.styles["Normal"]
normal.font.name = "Courier New"
normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Courier New")
normal.font.size = Pt(5.5)
normal.paragraph_format.left_indent = Cm(0)
normal.paragraph_format.right_indent = Cm(0)
normal.paragraph_format.first_line_indent = Cm(0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0
normal.paragraph_format.widow_control = False

# Remove the empty paragraph created with the document.
body = doc._element.body
for paragraph in list(doc.paragraphs):
    body.remove(paragraph._element)

for index, filename in enumerate(SOURCE_FILES):
    path = SOURCE_DIR / filename
    text = path.read_text(encoding="utf-8-sig")

    if index:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(""))

    separator = doc.add_paragraph()
    separator.paragraph_format.space_before = Pt(0)
    separator.paragraph_format.space_after = Pt(0)
    separator.paragraph_format.line_spacing = 1.0
    set_run_font(
        separator.add_run(
            f"// ===== BudgetTracker/{filename} ====="
        )
    )

    for raw_line in text.splitlines():
        # Word expands tab stops unpredictably in narrow columns.
        line = raw_line.expandtabs(4)
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0)
        paragraph.paragraph_format.right_indent = Cm(0)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.widow_control = False
        set_run_font(paragraph.add_run(line if line else " "))

doc.core_properties.title = "Листинг исходного кода BudgetTracker"
doc.core_properties.subject = "Исходный код проекта в три колонки"
doc.core_properties.author = ""
doc.core_properties.keywords = ""
doc.save(OUTPUT)

print(OUTPUT)
print(f"FILES={len(SOURCE_FILES)}")
print(f"SOURCE_CHARS={sum(len((SOURCE_DIR / f).read_text(encoding='utf-8-sig')) for f in SOURCE_FILES)}")
