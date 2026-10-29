from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = r"C:\Users\user\Desktop\windows-form\Раздел_3.1.1_Структура_программы.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2)
sec.bottom_margin = Cm(2)
sec.left_margin = Cm(3)
sec.right_margin = Cm(1.5)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(14)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.first_line_indent = Cm(1.25)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)

for name, size in [("Heading 1", 14), ("Heading 2", 14)]:
    st = styles[name]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.first_line_indent = Cm(0)
    st.paragraph_format.space_before = Pt(6)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True

caption = styles.add_style("CaptionCustom", WD_STYLE_TYPE.PARAGRAPH)
caption.font.name = "Times New Roman"
caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
caption.font.size = Pt(12)
caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_before = Pt(3)
caption.paragraph_format.space_after = Pt(6)
caption.paragraph_format.keep_with_next = True

def add_p(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.first_line_indent = Cm(0)
p.paragraph_format.space_after = Pt(6)
r = p.add_run("3.1.1 Структура программы (взаимосвязи функций, состав модулей и принципы объединения функций в модули)")
r.bold = True
r.font.name = "Times New Roman"
r._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
r._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
r.font.size = Pt(14)

add_p(
    "Разработанная программа «Финансовый помощник» представляет собой настольное приложение для учёта личных доходов и расходов. "
    "Программа реализована на языке C++/CLI с использованием технологии Windows Forms и платформы .NET Framework. "
    "В качестве постоянного хранилища применяется база данных PostgreSQL, в том числе допускается подключение к удалённому экземпляру Supabase. "
    "Архитектура построена по модульному принципу: пользовательский интерфейс, предметные данные, вычислительная логика и механизмы хранения разделены между самостоятельными компонентами."
)

add_p(
    "Точкой входа в приложение является модуль MyForm.cpp. В нём выполняется настройка визуальных стилей Windows Forms, создаётся объект главной формы MyForm и запускается цикл обработки сообщений. "
    "Главная форма описана в модуле MyForm.h. Она объединяет вкладки «Операции», «Личный профиль», «Категории и лимиты», «Регулярные платежи» и «Статистика», принимает действия пользователя и передаёт их соответствующим сервисам."
)

doc.add_paragraph("Состав программных модулей", style="Heading 2")

add_p(
    "Основу предметной области составляет модуль DomainModels.h. В нём определены классы финансовой операции, категории, лимита, регулярного платежа, личного профиля и общего снимка данных. "
    "Эти классы не содержат элементов интерфейса или кода подключения к базе данных и используются как единый формат обмена информацией между остальными частями программы."
)

add_p(
    "Изменение пользовательских данных сосредоточено в BudgetService.h. Сервис выполняет добавление, редактирование и удаление операций, управление категориями и их переименование. "
    "Проверка вводимых значений вынесена в InputValidator.h: данный модуль контролирует корректность денежных сумм и очищает текстовые значения. "
    "Такое разделение предотвращает дублирование проверок в обработчиках кнопок формы."
)

add_p(
    "Расчётные функции объединены по предметному назначению. FinanceService.h рассчитывает общий баланс, суммы доходов и расходов, а также группирует расходы по категориям. "
    "RecurringPaymentService.h отвечает за шаблоны ежемесячных платежей и создаёт операции, срок которых наступил. ForecastService.h формирует прогноз остатка средств до конца месяца с учётом текущих расходов и будущих регулярных платежей. "
    "ReportService.h создаёт текстовый отчёт за выбранный период."
)

add_p(
    "Работа с данными организована через интерфейс IDataStorage.h, который задаёт общие операции загрузки и сохранения. PostgreSqlStorage.h реализует данный контракт для PostgreSQL: создаёт необходимые таблицы, загружает записи текущего клиента и сохраняет изменения в рамках транзакции. "
    "SupabaseStorage.h предоставляет совместимый вариант доступа к удалённой базе Supabase. Выбор реализации выполняет StorageFactory.h, анализируя конфигурационный файл и переменную окружения. "
    "ClientIdentityService.h создаёт и хранит уникальный идентификатор установки, благодаря которому записи разных пользователей в общей базе не смешиваются."
)

doc.add_paragraph("Взаимосвязь компонентов программы", style="Heading 2")

doc.add_paragraph("Рисунок 1 — Логическая структура приложения", style="CaptionCustom")
table = doc.add_table(rows=4, cols=1)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
labels = [
    ("Пользовательский интерфейс — MyForm", "D9EAF7"),
    ("Прикладные сервисы — BudgetService, FinanceService,\nRecurringPaymentService, ForecastService, ReportService", "E8F1E8"),
    ("Предметные модели — DomainModels", "FFF2CC"),
    ("Слой хранения — IDataStorage → PostgreSqlStorage / SupabaseStorage", "F4CCCC"),
]
for i, (text, color) in enumerate(labels):
    cell = table.cell(i, 0)
    cell.width = Cm(15.7)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)
    shade(cell, color)
    p2 = cell.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.first_line_indent = Cm(0)
    p2.paragraph_format.line_spacing = 1.0
    p2.paragraph_format.space_after = Pt(0)
    run = p2.add_run(text)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(11)
    run.bold = True

add_p(
    "Связь модулей имеет направленный характер. Обработчики событий главной формы получают значения из полей и таблиц интерфейса, после чего используют InputValidator для проверки данных и вызывают методы прикладных сервисов. "
    "Сервисы работают с объектами DomainModels и не зависят от конкретных элементов Windows Forms. После успешного изменения форма передаёт общий снимок данных объекту IDataStorage, обновляет таблицы, баланс, диаграмму и прогноз."
)

add_p(
    "Например, при добавлении расхода процедура btnAdd_Click считывает сумму, категорию, дату и комментарий. После проверки выполняется расчёт уже потраченной суммы по категории и сравнение с установленным лимитом. "
    "Если лимит превышен, пользователю выводится предупреждение. После подтверждения BudgetService добавляет операцию, хранилище сохраняет изменения, а форма обновляет список операций и статистические показатели. "
    "Аналогичная последовательность применяется при редактировании и удалении записей."
)

add_p(
    "При запуске приложения StorageFactory создаёт объект хранилища, после чего загружается единый снимок данных. RecurringPaymentService проверяет регулярные платежи и при необходимости формирует новые операции. "
    "Затем MyForm распределяет загруженные сведения между элементами вкладок. При переходе на вкладку статистики повторно выполняются агрегация расходов, построение круговой диаграммы и расчёт прогноза. "
    "Таким образом, интерфейс всегда отображает актуальное состояние предметных данных."
)

doc.add_paragraph("Принципы объединения функций в модули", style="Heading 2")

add_p(
    "Функции объединены в модули по принципу единственной ответственности. MyForm отвечает только за отображение информации и обработку действий пользователя; DomainModels — за структуру данных; сервисные классы — за отдельные группы бизнес-правил; классы хранения — за обмен с базой данных. "
    "Модули связаны через понятные публичные методы и общие модели, поэтому изменение способа хранения не требует переработки расчётных функций."
)

add_p(
    "Дополнительно соблюдаются принципы слабой связанности и повторного использования. Расчёт баланса, прогнозирование и формирование отчёта не размещены непосредственно в обработчиках кнопок, поэтому могут вызываться из разных частей программы и отдельно тестироваться. "
    "Зависимость формы от интерфейса IDataStorage позволяет заменить локальный PostgreSQL удалённым Supabase без изменения пользовательских сценариев. Централизация проверки ввода, финансовых вычислений и сохранения данных уменьшает вероятность расхождений и упрощает дальнейшее развитие проекта."
)

add_p(
    "В результате структура программы образует последовательность «интерфейс — прикладные сервисы — предметные модели — хранилище». "
    "Такое построение соответствует назначению приложения, обеспечивает понятное распределение функций между файлами и позволяет независимо расширять интерфейс, финансовые алгоритмы и средства работы с базой данных."
)

# Footer with page number field.
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.paragraph_format.first_line_indent = Cm(0)
run = fp.add_run()
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
run._r.addnext(fld)

doc.save(OUT)
print(OUT)
